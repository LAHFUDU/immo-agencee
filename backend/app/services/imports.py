"""Bulk property import from CSV or Excel (.xlsx), with row validation and duplicate detection."""

import csv
import io
from decimal import Decimal, InvalidOperation

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Owner, Property, PropertyCategory, PropertyFeature
from app.models.enums import PropertyStatus, TransactionType
from app.services import audit, matching
from app.services.references import next_reference
from app.services.text import normalize_key

COLUMNS = [
    "reference", "type", "transaction", "titre", "description", "ville", "quartier", "adresse", "surface",
    "prix", "devise", "chambres", "salles_de_bain", "etage", "equipements", "statut", "publie",
    "latitude", "longitude", "proprietaire_nom", "proprietaire_telephone", "notes_internes",
]
REQUIRED = ["type", "transaction", "titre", "ville", "prix"]

ALIASES = {
    "ref": "reference", "property_type": "type", "type_bien": "type", "transaction_type": "transaction",
    "title": "titre", "city": "ville", "neighborhood": "quartier", "neighbourhood": "quartier", "address": "adresse",
    "surface_m2": "surface", "price": "prix", "currency": "devise", "bedrooms": "chambres", "bathrooms": "salles_de_bain",
    "floor": "etage", "features": "equipements", "status": "statut", "published": "publie",
    "owner_name": "proprietaire_nom", "owner_phone": "proprietaire_telephone", "internal_notes": "notes_internes",
}

TX_VALUES = {"vente": TransactionType.SALE, "sale": TransactionType.SALE, "achat": TransactionType.SALE,
             "location": TransactionType.RENT, "rent": TransactionType.RENT, "rental": TransactionType.RENT}
STATUS_VALUES = {
    "disponible": PropertyStatus.AVAILABLE, "available": PropertyStatus.AVAILABLE,
    "reserve": PropertyStatus.RESERVED, "reserved": PropertyStatus.RESERVED,
    "en negociation": PropertyStatus.UNDER_NEGOTIATION, "under negotiation": PropertyStatus.UNDER_NEGOTIATION,
    "vendu": PropertyStatus.SOLD, "sold": PropertyStatus.SOLD, "loue": PropertyStatus.RENTED, "rented": PropertyStatus.RENTED,
    "indisponible": PropertyStatus.UNAVAILABLE, "unavailable": PropertyStatus.UNAVAILABLE,
}
TRUE_VALUES = {"oui", "yes", "true", "1", "vrai", "x"}


class ImportFormatError(ValueError):
    pass


def _header_key(h) -> str:
    k = (normalize_key(str(h)) or "").replace(" ", "_")
    return ALIASES.get(k, k)


def read_rows(filename: str, data: bytes) -> list[dict]:
    name = filename.lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise ImportFormatError("Fichier Excel illisible.") from exc
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [_header_key(h) if h is not None else "" for h in rows[0]]
        out = []
        for r in rows[1:]:
            if r is None or all(v is None or str(v).strip() == "" for v in r):
                continue
            out.append({headers[i]: ("" if v is None else str(v).strip()) for i, v in enumerate(r) if i < len(headers)})
        return out
    if name.endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("cp1252")
        sample = text[:4096]
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        rows = list(reader)
        if not rows:
            return []
        headers = [_header_key(h) for h in rows[0]]
        return [
            {headers[i]: v.strip() for i, v in enumerate(r) if i < len(headers)}
            for r in rows[1:]
            if any(c.strip() for c in r)
        ]
    raise ImportFormatError("Format non supporté : utilisez un fichier .csv ou .xlsx.")


def _dec(v: str) -> Decimal:
    s = v.replace(" ", "").replace(" ", "")
    if "," in s and "." not in s:
        s = s.replace(",", ".")
    else:
        s = s.replace(",", "")
    return Decimal(s)


def _dup_key(type_, tx, city_key, nb_key, surface, price) -> tuple:
    return (type_, tx.value, city_key, nb_key or "", f"{Decimal(surface):.2f}" if surface is not None else "", f"{Decimal(price):.2f}")


def run_import(db: Session, filename: str, data: bytes, *, user_id, dry_run: bool, skip_duplicates: bool = True) -> dict:
    rows = read_rows(filename, data)
    categories = {c.code: c for c in db.scalars(select(PropertyCategory)).all()}
    cat_by_label = {normalize_key(c.label): c.code for c in categories.values()}
    features = {f.code: f for f in db.scalars(select(PropertyFeature)).all()}
    feat_by_label = {normalize_key(f.label): f.code for f in features.values()}

    existing_refs = set(db.scalars(select(Property.reference)).all())
    existing_keys = {
        _dup_key(p.property_type, p.transaction_type, p.city_key, p.neighborhood_key, p.surface, p.price): p.reference
        for p in db.scalars(select(Property).where(Property.status != PropertyStatus.ARCHIVED)).all()
    }

    errors: list[dict] = []
    duplicates: list[dict] = []
    valid: list[tuple[int, dict]] = []
    seen_refs: dict[str, int] = {}
    seen_keys: dict[tuple, int] = {}

    if rows and not any(k in rows[0] for k in REQUIRED):
        raise ImportFormatError("En-têtes de colonnes non reconnus. Téléchargez le modèle d'import.")

    for idx, raw in enumerate(rows, start=2):  # row 1 = headers
        row_errors = []

        def err(field, msg, row_errors=row_errors, idx=idx):
            row_errors.append({"row": idx, "field": field, "message": msg})

        for f in REQUIRED:
            if not raw.get(f):
                err(f, "Champ obligatoire manquant")
        item: dict = {}
        t = raw.get("type", "")
        code = t if t in categories else cat_by_label.get(normalize_key(t) or "")
        if t and not code:
            err("type", f"Type de bien inconnu : « {t} »")
        item["property_type"] = code
        tx = TX_VALUES.get(normalize_key(raw.get("transaction", "")) or "")
        if raw.get("transaction") and not tx:
            err("transaction", "Valeur attendue : vente ou location")
        item["transaction_type"] = tx
        item["title"] = raw.get("titre", "")[:200]
        item["description"] = raw.get("description") or None
        item["city"] = raw.get("ville", "")[:100]
        item["neighborhood"] = (raw.get("quartier") or None) and raw["quartier"][:120]
        item["address"] = raw.get("adresse") or None
        for src, dst, kind in [
            ("prix", "price", "dec"), ("surface", "surface", "dec"), ("chambres", "bedrooms", "int"),
            ("salles_de_bain", "bathrooms", "int"), ("etage", "floor", "int"),
            ("latitude", "latitude", "dec"), ("longitude", "longitude", "dec"),
        ]:
            v = raw.get(src, "")
            if v == "":
                item[dst] = None
                continue
            try:
                d = _dec(v)
                if kind == "int":
                    if d != d.to_integral_value():
                        raise InvalidOperation
                    item[dst] = int(d)
                else:
                    item[dst] = d
            except (InvalidOperation, ValueError):
                err(src, f"Nombre invalide : « {v} »")
                item[dst] = None
        if item.get("price") is not None and item["price"] < 0:
            err("prix", "Le prix doit être positif")
        if item.get("surface") is not None and item["surface"] <= 0:
            err("surface", "La surface doit être supérieure à 0")
        for f in ("bedrooms", "bathrooms"):
            if item.get(f) is not None and item[f] < 0:
                err(f, "Valeur négative")
        if item.get("latitude") is not None and not (-90 <= item["latitude"] <= 90):
            err("latitude", "Latitude hors limites")
        if item.get("longitude") is not None and not (-180 <= item["longitude"] <= 180):
            err("longitude", "Longitude hors limites")
        cur = (raw.get("devise") or "MAD").upper()
        if len(cur) != 3 or not cur.isalpha():
            err("devise", "Code devise sur 3 lettres attendu (ex. MAD, EUR)")
        item["currency"] = cur
        feats = []
        for part in [p.strip() for p in (raw.get("equipements") or "").replace(";", ",").split(",") if p.strip()]:
            fc = part if part in features else feat_by_label.get(normalize_key(part) or "")
            if not fc:
                err("equipements", f"Équipement inconnu : « {part} »")
            else:
                feats.append(fc)
        item["features"] = sorted(set(feats))
        st = raw.get("statut", "")
        status_val = STATUS_VALUES.get(normalize_key(st) or "", None) if st else PropertyStatus.AVAILABLE
        if status_val is None:
            err("statut", f"Statut inconnu : « {st} »")
        item["status"] = status_val or PropertyStatus.AVAILABLE
        item["is_published"] = (normalize_key(raw.get("publie", "")) or "") in TRUE_VALUES
        item["internal_notes"] = raw.get("notes_internes") or None
        item["owner_name"] = raw.get("proprietaire_nom") or None
        item["owner_phone"] = raw.get("proprietaire_telephone") or None
        ref = (raw.get("reference") or "").strip()
        item["reference"] = ref or None

        if row_errors:
            errors.extend(row_errors)
            continue

        city_key = normalize_key(item["city"])
        nb_key = normalize_key(item["neighborhood"])
        key = _dup_key(item["property_type"], item["transaction_type"], city_key, nb_key, item["surface"], item["price"])
        dup_reason = None
        if ref and ref in existing_refs:
            dup_reason = f"Référence {ref} déjà présente dans l'inventaire"
        elif ref and ref in seen_refs:
            dup_reason = f"Référence {ref} répétée (ligne {seen_refs[ref]})"
        elif key in existing_keys:
            dup_reason = f"Doublon probable du bien {existing_keys[key]} (même type, ville, quartier, surface et prix)"
        elif key in seen_keys:
            dup_reason = f"Doublon probable de la ligne {seen_keys[key]}"
        if dup_reason:
            duplicates.append({"row": idx, "reason": dup_reason})
            if skip_duplicates or (ref and (ref in existing_refs or ref in seen_refs)):
                continue
        if ref:
            seen_refs[ref] = idx
        seen_keys[key] = idx
        item["city_key"], item["neighborhood_key"] = city_key, nb_key
        valid.append((idx, item))

    imported: list[str] = []
    if not dry_run and valid:
        owners_cache: dict[tuple, Owner] = {}
        new_props = []
        for _, item in valid:
            owner = None
            if item["owner_name"]:
                ok = (item["owner_name"], item["owner_phone"])
                owner = owners_cache.get(ok)
                if owner is None:
                    owner = Owner(full_name=item["owner_name"], phone=item["owner_phone"], created_by_id=user_id)
                    db.add(owner)
                    owners_cache[ok] = owner
            p = Property(
                reference=item["reference"] or next_reference(db, "BIEN"),
                property_type=item["property_type"],
                transaction_type=item["transaction_type"],
                title=item["title"],
                description=item["description"],
                city=item["city"],
                city_key=item["city_key"],
                neighborhood=item["neighborhood"],
                neighborhood_key=item["neighborhood_key"],
                address=item["address"],
                surface=item["surface"],
                price=item["price"],
                currency=item["currency"],
                bedrooms=item["bedrooms"],
                bathrooms=item["bathrooms"],
                floor=item["floor"],
                features=item["features"],
                status=item["status"],
                is_published=item["is_published"],
                latitude=item["latitude"],
                longitude=item["longitude"],
                internal_notes=item["internal_notes"],
                owner=owner,
                responsible_id=user_id,
                created_by_id=user_id,
            )
            db.add(p)
            new_props.append(p)
        db.flush()
        for p in new_props:
            matching.refresh_for_property(db, p)
            imported.append(p.reference)
        audit.log(db, user_id=user_id, action="import", entity_type="property", label=filename,
                  changes={"imported": len(imported), "errors": len(errors), "duplicates": len(duplicates)})

    return {
        "total_rows": len(rows),
        "valid_rows": len(valid),
        "imported": len(imported),
        "imported_references": imported,
        "error_rows": len({e["row"] for e in errors}),
        "errors": errors,
        "duplicates": duplicates,
        "dry_run": dry_run,
    }


def template_csv() -> str:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(COLUMNS)
    w.writerow(["", "apartment", "vente", "Appartement lumineux 3 chambres", "Proche tramway", "Rabat", "Agdal",
                "12 rue Exemple", "120", "1150000", "MAD", "3", "2", "2", "parking,elevator", "disponible", "non",
                "", "", "M. Exemple", "0600000000", ""])
    return buf.getvalue()
